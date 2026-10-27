"""
Genera el informe de la EVA2 en Word usando la plantilla institucional de INACAP.

Uso:  python generar_informe_word.py
Entrada: plantilla/plantilla_inacap.docx
Salida:  informe/Informe_EVA2_Superstore.docx
"""
import re
import sys
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ / "src"))

import graficos  # noqa: E402
import limpieza  # noqa: E402
import textos    # noqa: E402

PLANTILLA = RAIZ / "plantilla" / "plantilla_inacap.docx"
SALIDA = RAIZ / "informe" / "Informe_EVA2_Superstore.docx"
CARPETA_FIG = RAIZ / "figuras"

AUTOR = "Diego Martínez González"
CARRERA = "Analista Programador"
ASIGNATURA = "Visualización de Datos (TIEL16) · EVA2"
PROFESOR = ""
URL_APP = "https://eva2-superstore-cf6parf9rmnwtmifgxebqn.streamlit.app/"

FUENTE = "Tahoma"
ROJO = RGBColor(0xEE, 0x00, 0x00)
ROJO_HEX = "EE0000"
GRIS_HEX = "BFBFBF"
FONDO_HEX = "F2F2F2"
VERDE_HEX = "E8F3EC"
ANCHO_UTIL_CM = 18.6          # carta (21,59 cm) menos márgenes de 1,5 cm


# ---------------------------------------------------------------------------
# Utilidades de formato (replican el formato directo de la plantilla)
# ---------------------------------------------------------------------------
def _fuente(run, size, color=None, bold=False, italic=False):
    run.font.name = FUENTE
    run._r.get_or_add_rPr().rFonts.set(qn("w:cs"), FUENTE)
    run.font.size = Pt(size)
    run.font.bold = bold or None
    run.font.italic = italic or None
    if color is not None:
        run.font.color.rgb = color


def _segmentos(texto):
    """Divide '**negrita**' y '__cursiva__' en (texto, negrita, cursiva)."""
    for p in re.split(r"(\*\*.+?\*\*|__.+?__)", texto):
        if p.startswith("**") and p.endswith("**"):
            yield p[2:-2], True, False
        elif p.startswith("__") and p.endswith("__"):
            yield p[2:-2], False, True
        elif p:
            yield p, False, False


def _escribir(par, texto, size, color=None, bold=False):
    for t, b, i in _segmentos(texto):
        _fuente(par.add_run(t), size, color, bold or b, i)


def _outline(par, nivel):
    o = OxmlElement("w:outlineLvl")
    o.set(qn("w:val"), str(nivel))
    par._p.get_or_add_pPr().append(o)


def _sombrear(celda, hex_):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    celda._tc.get_or_add_tcPr().append(sh)


def _insertar_tblpr(tblpr, nuevo, siguientes):
    """Inserta respetando el orden que exige el esquema de Word."""
    for hijo in tblpr:
        if hijo.tag in [qn(x) for x in siguientes]:
            hijo.addprevious(nuevo)
            return
    tblpr.append(nuevo)


def _bordes_tabla(tabla, color=GRIS_HEX):
    tblpr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    _insertar_tblpr(tblpr, b, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    m = OxmlElement("w:tblCellMar")
    for lado, w in (("top", 50), ("left", 80), ("bottom", 50), ("right", 80)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    _insertar_tblpr(tblpr, m, ("w:tblLook",))


def _enlace(par, url, size=11):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        f.set(qn(a), FUENTE)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    for e in (f, c, sz, u):
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = url
    r.append(t)
    h.append(r)
    par._p.append(h)


# ---------------------------------------------------------------------------
# Constructor del documento
# ---------------------------------------------------------------------------
class Informe:
    def __init__(self):
        self.doc = Document(str(PLANTILLA))
        self.n_fig = 0
        self._preparar_portada()

    def _preparar_portada(self):
        """Conserva portada y logo final de la plantilla; elimina el texto de ejemplo."""
        pars = self.doc.paragraphs
        titulo = next(p for p in pars if p.text.strip() == "TITULO PORTADA TRABAJO")
        subtitulo = next(p for p in pars if p.text.strip() == "Ejemplo de subtítulo")
        info = next(p for p in pars if p.text.startswith("NOMBRE:"))
        logo = next(p for p in pars if "<w:drawing" in p._p.xml)
        pos_info = next(i for i, p in enumerate(pars) if p._p is info._p)

        titulo.runs[0].text = textos.CONTEXTO["titulo"].upper()
        subtitulo.runs[0].text = "Informe de presentación de datos"
        valores = {"NOMBRE:": AUTOR, "CARRERA:": CARRERA, "ASIGNATURA:": ASIGNATURA,
                   "PROFESOR:": PROFESOR, "FECHA:": f"{date.today():%d-%m-%Y}"}
        for r in info.runs:
            clave = r.text.strip()
            if clave in valores:
                r.text = f"{clave} {valores[clave]}".rstrip()

        for i, p in enumerate(pars):
            if i > pos_info and p._p is not logo._p:
                p._p.getparent().remove(p._p)
        self.logo = logo._p
        self.logo.getparent().remove(self.logo)
        self.salto_pagina()

    def cerrar(self):
        """Contraportada: el logo institucional va en una página propia."""
        self.salto_pagina()
        self.doc.element.body.sectPr.addprevious(self.logo)

    # -- bloques -----------------------------------------------------------------
    def salto_pagina(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, numero, texto):
        p = self.doc.add_paragraph()
        p.alignment = 1
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        _escribir(p, f"{numero} {texto}", 18, ROJO)
        _outline(p, 0)

    def h2(self, numero, texto):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        _escribir(p, f"{numero} {texto}", 14, ROJO)
        _outline(p, 1)

    def texto(self, texto, size=11, color=None, despues=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(despues)
        _escribir(p, texto, size, color)
        return p

    def nota(self, texto):
        return self.texto(texto, size=9, color=RGBColor(0x59, 0x59, 0x59))

    def vineta(self, texto):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(4)
        _escribir(p, "‐ " + texto, 11)

    def caja(self, texto, fondo=FONDO_HEX):
        t = self.doc.add_table(rows=1, cols=1)
        _bordes_tabla(t, fondo)
        c = t.cell(0, 0)
        _sombrear(c, fondo)
        c.width = Cm(ANCHO_UTIL_CM)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _escribir(p, texto, 10.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def tabla(self, encabezado, filas, anchos_cm, size=8):
        t = self.doc.add_table(rows=1, cols=len(encabezado))
        t.autofit = False
        _bordes_tabla(t)
        for i, (txt, w) in enumerate(zip(encabezado, anchos_cm)):
            c = t.rows[0].cells[i]
            c.width = Cm(w)
            _sombrear(c, ROJO_HEX)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _escribir(p, txt, size, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for fila in filas:
            celdas = t.add_row().cells
            for i, (txt, w) in enumerate(zip(fila, anchos_cm)):
                celdas[i].width = Cm(w)
                p = celdas[i].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                _escribir(p, str(txt), size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figura(self, fig, nombre, titulo, ancho_cm=16.5):
        CARPETA_FIG.mkdir(exist_ok=True)
        ruta = CARPETA_FIG / f"{nombre}.png"
        fig.savefig(ruta, bbox_inches="tight")
        plt.close(fig)
        self.n_fig += 1
        p = self.doc.add_paragraph()
        p.alignment = 1
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
        c = self.doc.add_paragraph()
        c.alignment = 1
        _escribir(c, f"Figura {self.n_fig}. {titulo}", 9, RGBColor(0x59, 0x59, 0x59))

    def guardar(self):
        SALIDA.parent.mkdir(exist_ok=True)
        cp = self.doc.core_properties
        cp.title = textos.CONTEXTO["titulo"]
        cp.author = AUTOR
        cp.subject = ASIGNATURA
        self.doc.save(str(SALIDA))
        print(f"Informe generado: {SALIDA}")


def _num(n):
    return f"{int(n):,}".replace(",", ".")


# ---------------------------------------------------------------------------
# Contenido (mismas fuentes que la app: src/limpieza, src/graficos, src/textos)
# ---------------------------------------------------------------------------
def construir():
    graficos.configurar_estilo()
    crudo = limpieza.cargar_crudo()
    df, bitacora, resumen = limpieza.limpiar()
    fichas = textos.ficha_visualizaciones(df)
    plt.close("all")
    atipicos = limpieza.efecto_eliminar_atipicos(df)
    ctx = textos.CONTEXTO
    w = Informe()

    # 1. Contexto -----------------------------------------------------------------
    w.h1("1", "Contexto")
    w.texto(ctx["descripcion"])
    w.caja(f"**Pregunta central.** {ctx['pregunta_central']}")
    p = w.texto(f"**Dataset:** {ctx['dataset']} ({ctx['fuente']}). {_num(resumen['filas_iniciales'])} líneas de "
                f"pedido, 21 columnas, periodo {resumen['periodo']}. **Herramienta:** Python (pandas, matplotlib, "
                "seaborn) y una app en Streamlit: ")
    _enlace(p, URL_APP)
    w.h2("1.1", "Preguntas que responde el informe")
    w.tabla(["#", "Tipo (tabla U1.4)", "Pregunta de negocio", "Gráfico"],
            [[i, f["tipo"], f["pregunta"], f["grafico"]] for i, f in enumerate(fichas, 1)],
            [0.8, 4.0, 8.0, 5.8], size=9)

    # 2. Preparación de datos -------------------------------------------------------
    w.salto_pagina()
    w.h1("2", "Preparación de datos")
    w.texto("El dataset se revisó en tres dimensiones: **integridad** (registros completos y sin duplicar), "
            "**calidad** (valores correctos y útiles) y **homologación** (una misma cosa escrita siempre igual). "
            f"No hay valores nulos ({int(crudo.isna().sum().sum())}), pero sí problemas menos visibles. "
            f"Se documentaron {len(bitacora)} decisiones; la limpieza deja {_num(resumen['filas_finales'])} filas "
            f"de las {_num(resumen['filas_iniciales'])} originales.")
    filas = [[r["Dimensión"], r["Problema detectado"], _num(r["Filas afectadas"]), r["Decisión"],
              f"{r['Justificación']} __{r['Implicancia']}__"] for _, r in bitacora.iterrows()]
    w.tabla(["Dimensión", "Problema detectado", "Filas", "Decisión", "Justificación e implicancia"], filas,
            [2.2, 4.4, 1.3, 4.0, 6.7])
    w.nota("Las filas afectadas cuentan registros revisados o transformados, no eliminados. Solo se eliminó "
           "1 fila (el duplicado exacto). Ninguna región pierde representación.")

    # 3. Visualizaciones --------------------------------------------------------------
    w.salto_pagina()
    w.h1("3", "Visualizaciones e interpretación de negocio")
    w.texto("Cada gráfico responde una pregunta distinta y termina en una decisión. Todos se construyeron con "
            "matplotlib y seaborn con una configuración común: fondo claro, grilla tenue, títulos que enuncian "
            "el hallazgo, montos en dólares con separador de miles chileno y una paleta validada para "
            "personas con daltonismo.")
    for i, f in enumerate(fichas, 1):
        if i > 1:
            w.salto_pagina()
        fig, _ = f["funcion"](df)
        w.h2(f"3.{i}", f"{f['tipo']}: {f['pregunta']}")
        w.figura(fig, f["id"], f["grafico"], 15.0 if f["id"] == "g5" else 16.5)
        w.texto(f"**Qué muestra.** {f['lectura']}")
        w.caja(f"**Decisión que apoya.** {f['decision']}", VERDE_HEX)
        w.nota(f"**Por qué este gráfico.** {f['por_que']}")
        w.nota(f"**Configuración.** {f['configuracion']}")

    # 4. Ética -------------------------------------------------------------------------
    w.salto_pagina()
    w.h1("4", "Decisiones con implicancia ética")
    w.texto("Limpiar y graficar obliga a decidir qué se muestra y qué no. Estas son las decisiones que podían "
            "cambiar las conclusiones.")
    for i, e in enumerate(textos.reflexion_etica(df, atipicos, resumen), 1):
        w.h2(f"4.{i}", e["decision"])
        w.texto(f"**Situación.** {e['contexto']}")
        w.texto(f"**Consecuencia.** {e['consecuencia']}")
        if i == 1:
            filas = [[reg, _num(r["Filas totales"]),
                      f"{int(r['Filas que se eliminarían'])} ({graficos._pct(r['% de filas perdidas'])})",
                      graficos._pct(r["Margen real"], 1), graficos._pct(r["Margen sin atípicos"], 1),
                      f"{r['Diferencia (pp)']:+.1f} pp".replace(".", ",")] for reg, r in atipicos.iterrows()]
            w.tabla(["Región", "Filas", "Filas eliminadas", "Margen real", "Margen sin atípicos", "Diferencia"],
                    filas, [2.8, 2.2, 3.8, 3.0, 3.8, 3.0], size=9)

    # 5. Conclusiones --------------------------------------------------------------------
    w.h1("5", "Conclusiones")
    for c in textos.conclusiones(df):
        w.vineta(c)
    w.nota("Código fuente y app: el repositorio contiene src/limpieza.py (limpieza y bitácora), "
           "src/graficos.py (visualizaciones), src/textos.py (interpretaciones) y app.py (Streamlit). "
           "Este informe se genera con generar_informe_word.py desde el mismo código.")

    w.cerrar()
    w.guardar()


if __name__ == "__main__":
    construir()
